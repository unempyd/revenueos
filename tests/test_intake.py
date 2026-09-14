"""Document intake, against real files built at test time — no binary fixtures in the repo.

The PDF is written by a small writer in this file (a valid PDF is a handful of objects and a
content stream); the DOCX, XLSX and PPTX are written by the same libraries `intake.py` reads
them back with, which is the only way to be sure the round trip is real.
"""
from __future__ import annotations

import json
from pathlib import Path

import pytest

from revenueos.intake import MAX_BYTES, Document, read_document
from revenueos.store import Store
from revenueos.workers import execute_action, run_worker


# ── fixture builders ─────────────────────────────────────────────────────────
def write_pdf(path: Path, pages: list[list[str]], title: str = "", author: str = "") -> Path:
    """A real, minimal PDF: Helvetica text at 12pt, one text block per page. A page given an
    empty line list gets a content stream with no text-showing operator — i.e. a scan."""
    def esc(s: str) -> bytes:
        return s.replace("\\", r"\\").replace("(", r"\(").replace(")", r"\)").encode("latin-1", "replace")

    n = len(pages)
    objs: list[bytes] = [
        b"<< /Type /Catalog /Pages 2 0 R >>",
        f"<< /Type /Pages /Kids [{' '.join(f'{5 + 2 * i} 0 R' for i in range(n))}] /Count {n} >>".encode(),
        b"<< /Type /Font /Subtype /Type1 /BaseFont /Helvetica >>",
    ]
    for i, lines in enumerate(pages):
        body = b"BT /F1 12 Tf 72 720 Td 16 TL\n"
        for ln in lines:
            body += b"(" + esc(ln) + b") Tj T*\n"
        body += b"ET"
        objs.append(b"<< /Length " + str(len(body)).encode() + b" >>\nstream\n" + body + b"\nendstream")
        objs.append((f"<< /Type /Page /Parent 2 0 R /MediaBox [0 0 612 792] /Contents {4 + 2 * i} 0 R "
                     f"/Resources << /Font << /F1 3 0 R >> >> >>").encode())
    info = None
    if title or author:
        objs.append(b"<< /Title (" + esc(title) + b") /Author (" + esc(author) + b") >>")
        info = len(objs)
    out = bytearray(b"%PDF-1.4\n")
    offsets = []
    for i, o in enumerate(objs, start=1):
        offsets.append(len(out))
        out += f"{i} 0 obj\n".encode() + o + b"\nendobj\n"
    xref = len(out)
    out += f"xref\n0 {len(objs) + 1}\n".encode() + b"0000000000 65535 f \n"
    for off in offsets:
        out += f"{off:010d} 00000 n \n".encode()
    trailer = f"trailer\n<< /Size {len(objs) + 1} /Root 1 0 R" + (f" /Info {info} 0 R" if info else "")
    out += (trailer + f" >>\nstartxref\n{xref}\n%%EOF\n").encode()
    path.write_bytes(bytes(out))
    return path


def write_docx(path: Path) -> Path:
    import docx

    d = docx.Document()
    d.core_properties.title = "Acme Recall — Brand and Positioning Guide"
    d.core_properties.author = "Sam Rivera"
    d.add_heading("Acme Recall — Brand and Positioning Guide", level=1)
    d.add_paragraph("Acme Recall fills empty dental chairs with two-way SMS rebooking.")
    d.add_heading("Who we sell to", level=2)
    d.add_paragraph("Practice managers at 2-10 chair dental clinics in the United States.")
    d.add_heading("Words we never use", level=2)
    d.add_paragraph("revolutionary", style="List Bullet")
    d.add_paragraph("cutting-edge", style="List Bullet")
    table = d.add_table(rows=3, cols=3)
    for r, row in enumerate([["Plan", "Price", "Seats"], ["Basic", "$149/mo", "1 location"], ["Pro", "$399/mo", "5 locations"]]):
        for c, value in enumerate(row):
            table.cell(r, c).text = value
    d.save(str(path))
    return path


def write_xlsx(path: Path) -> Path:
    import openpyxl

    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Price list"
    for row in [["SKU", "Plan", "Monthly", "Annual"], ["ACM-B", "Basic", 149, 1490],
                ["ACM-P", "Pro", 399, 3990], ["ACM-E", "Enterprise", 899, 8990]]:
        ws.append(row)
    second = wb.create_sheet("Discounts")
    for row in [["Term", "Discount"], ["Annual", "2 months free"], ["3 year", "20%"]]:
        second.append(row)
    wb.properties.creator = "Sam Rivera"
    wb.save(str(path))
    return path


def write_pptx(path: Path) -> Path:
    from pptx import Presentation

    prs = Presentation()
    slide = prs.slides.add_slide(prs.slide_layouts[1])
    slide.shapes.title.text = "Q3 competitor review"
    slide.placeholders[1].text = "Weave leads on bundling. NexHealth leads on PMS integrations."
    second = prs.slides.add_slide(prs.slide_layouts[1])
    second.shapes.title.text = "Where we win"
    second.placeholders[1].text = "Two-way SMS rebooking; installs in an afternoon."
    prs.core_properties.author = "Sam Rivera"
    prs.save(str(path))
    return path


# ── per-format round trips ───────────────────────────────────────────────────
def test_pdf_round_trips(tmp_path: Path):
    p = write_pdf(tmp_path / "price-list.pdf",
                  [["Acme Recall price list 2026", "Basic plan $149 per month", "Pro plan $399 per month"],
                   ["Terms", "Net 30. Annual billing gets two months free."]],
                  title="Acme Recall price list", author="Sam Rivera")
    doc = read_document(p)
    assert doc.ok, doc.error
    assert doc.format == "pdf" and doc.pages == 2 and doc.unit == "page"
    assert "Basic plan $149 per month" in doc.text
    assert "Net 30" in doc.text
    assert "## Page 2" in doc.text          # page structure is preserved as headings
    assert doc.metadata["title"] == "Acme Recall price list"
    assert doc.metadata["author"] == "Sam Rivera"
    assert doc.sha256 and doc.size_bytes == p.stat().st_size


def test_docx_round_trips_with_headings_and_table(tmp_path: Path):
    doc = read_document(write_docx(tmp_path / "brand.docx"))
    assert doc.ok, doc.error
    assert doc.format == "docx"
    assert "# Acme Recall — Brand and Positioning Guide" in doc.text
    assert "## Who we sell to" in doc.text                      # heading level survives
    assert "- revolutionary" in doc.text                        # list style survives
    assert doc.metadata["author"] == "Sam Rivera"
    assert len(doc.tables) == 1
    assert doc.tables[0].rows[0] == ["Plan", "Price", "Seats"]
    assert ["Pro", "$399/mo", "5 locations"] in doc.tables[0].rows
    # the table stays where the author put it — after the bullets, not shunted to the end
    assert doc.text.index("- cutting-edge") < doc.text.index("| Plan | Price | Seats |")
    # .docx carries no real page count: whatever is in docProps/app.xml is what the saving
    # application recorded, and the note says exactly that rather than passing it off as measured
    assert any("page count" in n for n in doc.notes)


def test_xlsx_round_trips_every_sheet_as_rows(tmp_path: Path):
    doc = read_document(write_xlsx(tmp_path / "prices.xlsx"))
    assert doc.ok, doc.error
    assert doc.format == "xlsx" and doc.pages == 2 and doc.unit == "sheet"
    assert [t.name for t in doc.tables] == ["Price list", "Discounts"]
    assert doc.tables[0].rows[0] == ["SKU", "Plan", "Monthly", "Annual"]
    assert ["ACM-P", "Pro", "399", "3990"] in doc.tables[0].rows
    assert "## Sheet: Discounts" in doc.text


def test_pptx_round_trips(tmp_path: Path):
    doc = read_document(write_pptx(tmp_path / "competitors.pptx"))
    assert doc.ok, doc.error
    assert doc.format == "pptx" and doc.pages == 2 and doc.unit == "slide"
    assert "## Slide 1: Q3 competitor review" in doc.text
    assert "NexHealth leads on PMS integrations" in doc.text


def test_csv_and_text_and_rtf(tmp_path: Path):
    csv_path = tmp_path / "leads.csv"
    csv_path.write_text("company,website,plan\nAcme Dental,acmedental.example,Pro\nBright Smiles,bright.example,Basic\n")
    csv_doc = read_document(csv_path)
    assert csv_doc.ok and csv_doc.tables[0].rows[0] == ["company", "website", "plan"]
    assert ["Bright Smiles", "bright.example", "Basic"] in csv_doc.tables[0].rows

    md = tmp_path / "plan.md"
    md.write_text("# Business plan\n\nGrow to 200 clinics by Q4.\n")
    assert read_document(md).text.startswith("# Business plan")

    rtf = tmp_path / "letter.rtf"
    rtf.write_text(r"{\rtf1\ansi Acme Recall is \b $149 per month\b0 .}")
    rtf_doc = read_document(rtf)
    assert rtf_doc.ok and "149 per month" in rtf_doc.text


# ── failure is a sentence, never a crash and never a partial lie ─────────────
def test_corrupt_pdf_returns_ok_false_with_a_sentence(tmp_path: Path):
    p = tmp_path / "broken.pdf"
    p.write_bytes(b"%PDF-1.4\nthis is not really a PDF at all\n")
    doc = read_document(p)
    assert doc.ok is False
    assert doc.text == "" and doc.tables == []
    assert doc.error and doc.error.endswith(".") and "could not be read" in doc.error


def test_corrupt_docx_is_named_not_guessed(tmp_path: Path):
    p = tmp_path / "broken.docx"
    p.write_bytes(b"not a zip archive at all, just bytes")
    doc = read_document(p)
    assert doc.ok is False and "not a zip archive" in doc.error


def test_oversized_file_is_refused_before_parsing(tmp_path: Path):
    big = tmp_path / "huge.pdf"
    with open(big, "wb") as fh:          # sparse: no 26 MB actually written to disk
        fh.truncate(MAX_BYTES + 1)
    doc = read_document(big)
    assert doc.ok is False and "the limit is" in doc.error
    assert doc.sha256 == ""              # it was refused before the bytes were even hashed


def test_text_free_pdf_is_reported_as_a_scan(tmp_path: Path):
    p = write_pdf(tmp_path / "scan.pdf", [[], []])   # two real pages, no text-showing operator
    doc = read_document(p)
    assert doc.ok is False
    assert doc.pages == 2                            # the pages are real; the text is not there
    assert "this looks like a scan" in doc.error and "OCR" in doc.error


def test_legacy_and_unsupported_formats_are_refused_by_name(tmp_path: Path):
    old = tmp_path / "plan.doc"
    old.write_bytes(b"\xd0\xcf\x11\xe0" + b"\x00" * 64)
    assert ".docx" in read_document(old).error
    weird = tmp_path / "thing.xyz"
    weird.write_text("hello")
    assert "cannot read .xyz" in read_document(weird).error
    assert read_document(tmp_path / "absent.pdf").error == "there is no file at that path."


def test_text_is_truncated_with_an_explicit_marker(tmp_path: Path):
    p = tmp_path / "long.md"
    p.write_text("Acme Recall. " * 5000)
    doc = read_document(p, max_chars=2000)
    assert doc.ok and doc.truncated and len(doc.text) < 2200
    assert "[truncated by RevenueOS" in doc.text
    assert any("truncated" in n for n in doc.notes)


def test_zip_bomb_is_refused(tmp_path: Path, monkeypatch: pytest.MonkeyPatch):
    import revenueos.intake as intake

    p = write_xlsx(tmp_path / "prices.xlsx")
    monkeypatch.setattr(intake, "MAX_UNCOMPRESSED", 1024)     # this workbook expands past 1 KB
    doc = read_document(p)
    assert doc.ok is False and "will not open it" in doc.error


# ── the worker in the loop ───────────────────────────────────────────────────
PROPOSAL = json.dumps({
    "summary": "A price list: Basic $149/mo, Pro $399/mo, Enterprise $899/mo, annual billing two months free.",
    "proposals": [
        {"file": "offer.md", "heading": "Core Offer",
         "proposed": "Acme Recall — Basic $149/month, Pro $399/month, Enterprise $899/month per location.",
         "why": "The canon names one price; the price list has three tiers.",
         "quote": "Pro plan $399 per month"},
        {"file": "offer.md", "heading": "A Heading That Does Not Exist",
         "proposed": "nonsense", "why": "invented", "quote": "invented"},
    ],
})


class FakeLLM:
    """Stands in for the one LLM client; records what it was asked."""

    def __init__(self, reply: str = PROPOSAL) -> None:
        self.reply = reply
        self.prompts: list[str] = []

    def complete_sync(self, messages, *, temperature=None, max_tokens=4096) -> str:
        self.prompts.append("\n".join(m.content for m in messages))
        return self.reply


def _drop(ws) -> list[Path]:
    inbox = ws.documents_inbox
    return [write_pdf(inbox / "price-list.pdf",
                      [["Acme Recall price list 2026", "Basic plan $149 per month", "Pro plan $399 per month"]],
                      title="Acme Recall price list"),
            write_docx(inbox / "brand.docx"),
            write_xlsx(inbox / "prices.xlsx")]


def test_worker_reads_every_dropped_document_and_is_idempotent(workspace, store: Store, onboarded):
    _drop(workspace)
    first = run_worker("intake", workspace, store, onboarded, None)
    assert first.ok, first.error
    assert first.details["read"] == 3 and first.details["failed"] == 0
    assert len(store.list_documents()) == 3
    assert {d["format"] for d in store.list_documents()} == {"pdf", "docx", "xlsx"}
    written = sorted(p.name for p in workspace.documents.iterdir())
    assert len(written) == 3 and all(n.endswith(".md") for n in written)

    second = run_worker("intake", workspace, store, onboarded, None)
    assert second.details["read"] == 0 and second.details["skipped"] == 3
    assert second.actions_created == 0
    assert len(store.list_documents()) == 3
    assert sorted(p.name for p in workspace.documents.iterdir()) == written


def test_worker_without_a_model_extracts_but_does_not_summarise(workspace, store: Store, onboarded):
    _drop(workspace)
    result = run_worker("intake", workspace, store, onboarded, None)
    assert result.ok and result.actions_created == 0
    assert "needs a model credential" in result.summary
    assert all(d["summary"] is None for d in store.list_documents())
    assert all(d["text_path"] and d["chars"] > 0 for d in store.list_documents())
    assert store.list_actions("pending", "correction") == []


def test_unreadable_document_is_reported_and_not_retried(workspace, store: Store, onboarded):
    (workspace.documents_inbox / "broken.pdf").write_bytes(b"%PDF-1.4\nnot a pdf\n")
    result = run_worker("intake", workspace, store, onboarded, None)
    assert result.ok and result.details["failed"] == 1 and result.details["read"] == 0
    row = store.list_documents()[0]
    assert row["error"] and row["text_path"] is None
    assert run_worker("intake", workspace, store, onboarded, None).details["skipped"] == 1


def _canon_snapshot(ws) -> dict[str, str]:
    return {p.name: p.read_text() for p in sorted(ws.company_context.glob("*.md"))}


def test_a_document_proposes_a_pending_correction_and_never_edits_the_canon(workspace, store: Store, onboarded):
    _drop(workspace)
    before = _canon_snapshot(workspace)
    llm = FakeLLM()
    result = run_worker("intake", workspace, store, onboarded, llm)

    assert result.ok, result.error
    assert _canon_snapshot(workspace) == before, "the worker must never write to company-context/"

    actions = store.list_actions("pending", "correction")
    assert actions, "a contradicted heading should become a pending correction"
    assert all(a["status"] == "pending" for a in actions)
    # the invented heading in the model's reply is dropped, not repaired
    assert all(a["context"]["canon_heading"] == "Core Offer" for a in actions)
    assert all("$399" in a["content"] for a in actions)
    assert all(a["context"]["executor"] == "record_correction" for a in actions)
    assert store.list_documents()[0]["summary"].startswith("A price list")
    # the model saw the canon and the extracted text, one prompt per document — never the file bytes
    assert len(llm.prompts) == 3
    everything = "\n".join(llm.prompts)
    assert "Acme Scheduling" in everything                  # the canon RevenueOS already holds
    assert "Basic plan $149 per month" in everything        # the PDF's text
    assert "| Pro | $399/mo | 5 locations |" in everything  # the DOCX's table, as rows

    # running again proposes nothing new: the same document hash, the same dedupe key
    again = run_worker("intake", workspace, store, onboarded, FakeLLM())
    assert again.actions_created == 0
    assert len(store.list_actions("pending", "correction")) == len(actions)


def test_approving_a_correction_writes_to_the_learning_loop_only(workspace, store: Store, onboarded):
    _drop(workspace)
    run_worker("intake", workspace, store, onboarded, FakeLLM())
    action = store.list_actions("pending", "correction")[0]
    before = _canon_snapshot(workspace)

    outcome = execute_action(workspace, store, onboarded, None, action)

    assert "learning-loop/CORRECTIONS.md" in outcome
    assert _canon_snapshot(workspace) == before
    corrections = workspace.corrections.read_text()
    assert "offer.md → Core Offer" in corrections and "$399/month" in corrections


def test_empty_drop_folder_says_where_to_put_documents(workspace, store: Store, onboarded):
    result = run_worker("intake", workspace, store, onboarded, None)
    assert result.ok and result.actions_created == 0
    assert "data/inbox/documents" in result.summary


def test_document_as_markdown_carries_its_provenance(tmp_path: Path):
    doc: Document = read_document(write_xlsx(tmp_path / "prices.xlsx"))
    md = doc.as_markdown()
    assert md.startswith("# ")
    assert "prices.xlsx" in md and "2 sheets" in md
