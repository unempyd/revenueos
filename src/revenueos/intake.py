"""INTAKE — turn the business's own documents into text RevenueOS can reason over.

A business already has its price list, its brand guide, its business plan and last quarter's
competitor report.  Until now nothing in RevenueOS could open any of them, so all of that had
to be retyped into the onboarding questionnaire.  This module reads them.

**Extraction is entirely local.**  No document byte ever leaves this machine: every format is
parsed in-process by a pure-Python library (no API call, no upload, no conversion service, and
no shelling out to a binary such as `pdftotext`, `libreoffice` or `antiword`, which a customer
machine may not have).  Only the *extracted text* — and only if the workspace has a model
credential — is later sent to a model by `workers/intake.py`, exactly like any other worker
prompt.

    read_document(path) -> Document      one file in, structured text out; never raises

Supported (extension → parser, all MIT/BSD/Apache):

    .pdf                 pypdf            (BSD-3-Clause)
    .docx                python-docx      (MIT)
    .xlsx .xlsm          openpyxl         (MIT)
    .pptx                python-pptx      (MIT)
    .csv .tsv            stdlib csv
    .txt .md .markdown   stdlib
    .rtf                 striprtf         (BSD-3-Clause)

Anything else — .doc, .ppt, .xls, .pages, .odt, .epub — is refused by name rather than
guessed at.  A document that cannot be parsed comes back `ok=False` with one plain sentence;
it never raises and it never returns half a document as if it were the whole one.

Limits are hard, because a document can be hostile or simply enormous (`Document.notes` says
whenever one bit):

    MAX_BYTES            25 MiB on disk
    MAX_UNCOMPRESSED     300 MiB expanded, for the zip-based Office formats (zip-bomb guard)
    MAX_UNITS            300 pages / sheets / slides
    MAX_CHARS            400,000 characters of text, then an explicit truncation marker
    TIME_BUDGET_S        60 s of wall clock, checked between pages/sheets/slides

The wall-clock guard is checked between units, which is the honest bound a pure-Python parser
allows: it cannot interrupt a single pathological page mid-parse, only stop before the next one.
"""
from __future__ import annotations

import csv
import hashlib
import io
import logging
import re
import time
import zipfile
from dataclasses import dataclass, field
from pathlib import Path
from typing import Any

MAX_BYTES = 25 * 1024 * 1024
MAX_UNCOMPRESSED = 300 * 1024 * 1024
MAX_UNITS = 300
MAX_CHARS = 400_000
MAX_TABLE_ROWS = 500
MAX_TABLE_COLS = 60
MAX_CELL_CHARS = 500
TIME_BUDGET_S = 60.0

TRUNCATION_MARKER = "\n\n[truncated by RevenueOS: the document is longer than {limit:,} characters]"

# extension → (format name, unit word). The unit is what `pages` counts, or None when the
# format carries no such count at all.
FORMATS: dict[str, tuple[str, str | None]] = {
    ".pdf": ("pdf", "page"),
    ".docx": ("docx", "page"),
    ".xlsx": ("xlsx", "sheet"),
    ".xlsm": ("xlsx", "sheet"),
    ".pptx": ("pptx", "slide"),
    ".csv": ("csv", None),
    ".tsv": ("csv", None),
    ".txt": ("text", None),
    ".md": ("text", None),
    ".markdown": ("text", None),
    ".rtf": ("rtf", None),
}
SUPPORTED_EXTENSIONS = tuple(FORMATS)

# Formats a business will genuinely hand over that this module deliberately does not guess at.
REFUSED: dict[str, str] = {
    ".doc": "the old binary Word format. Open it in Word or Pages and save it as .docx, then drop that in.",
    ".xls": "the old binary Excel format. Open it and save it as .xlsx, then drop that in.",
    ".ppt": "the old binary PowerPoint format. Open it and save it as .pptx, then drop that in.",
    ".pages": "an Apple Pages bundle. Export it as PDF or Word (.docx) and drop that in.",
    ".key": "an Apple Keynote bundle. Export it as PDF or PowerPoint (.pptx) and drop that in.",
    ".numbers": "an Apple Numbers bundle. Export it as Excel (.xlsx) or CSV and drop that in.",
    ".odt": "OpenDocument text is not supported yet. Save it as .docx or PDF and drop that in.",
    ".ods": "OpenDocument spreadsheet is not supported yet. Save it as .xlsx or CSV and drop that in.",
    ".odp": "OpenDocument presentation is not supported yet. Save it as .pptx or PDF and drop that in.",
    ".epub": "EPUB is not supported yet. Export the chapters you want as PDF and drop that in.",
}

SCAN_SENTENCE = ("no extractable text: this looks like a scan (a picture of the pages rather than text). "
                 "Run it through OCR first, or send the original file it was printed from.")


@dataclass
class Table:
    """One table exactly as the format carried it. `rows[0]` is the header row when the format
    marks one (spreadsheet first row, Word/PowerPoint first table row); nothing is inferred."""

    name: str
    rows: list[list[str]] = field(default_factory=list)
    truncated: bool = False

    def as_markdown(self) -> str:
        if not self.rows:
            return ""
        head, *body = self.rows
        width = len(head)
        out = ["| " + " | ".join(head) + " |", "|" + "|".join([" --- "] * width) + "|"]
        for r in body:
            cells = (r + [""] * width)[:width]
            out.append("| " + " | ".join(cells) + " |")
        if self.truncated:
            out.append(f"| … {MAX_TABLE_ROWS} rows shown; the table is longer |" + " |" * (width - 1))
        return "\n".join(out)

    def as_json(self) -> dict[str, Any]:
        return {"name": self.name, "rows": self.rows, "truncated": self.truncated}


@dataclass
class Document:
    """The result of reading one file. `ok=False` always carries `error`, one plain sentence."""

    path: Path
    format: str
    ok: bool = False
    text: str = ""
    pages: int | None = None          # pages / sheets / slides, or None when the format has no such count
    unit: str | None = None           # "page" | "sheet" | "slide" | None
    tables: list[Table] = field(default_factory=list)
    metadata: dict[str, str] = field(default_factory=dict)   # title / author / created, where present
    sha256: str = ""
    size_bytes: int = 0
    truncated: bool = False
    notes: list[str] = field(default_factory=list)
    error: str | None = None

    @property
    def title(self) -> str:
        return (self.metadata.get("title") or "").strip() or self.path.name

    @property
    def chars(self) -> int:
        return len(self.text)

    def describe(self) -> str:
        """One line for a CLI or a run summary."""
        if not self.ok:
            return f"{self.path.name}: {self.error}"
        bits = [self.format]
        if self.pages is not None:
            bits.append(f"{self.pages} {self.unit}{'s' if self.pages != 1 else ''}")
        bits.append(f"{self.chars:,} chars")
        if self.tables:
            bits.append(f"{len(self.tables)} table{'s' if len(self.tables) != 1 else ''}")
        return f"{self.path.name}: " + ", ".join(bits)

    def as_json(self) -> dict[str, Any]:
        return {"path": str(self.path), "format": self.format, "ok": self.ok, "pages": self.pages,
                "unit": self.unit, "chars": self.chars, "title": self.title, "sha256": self.sha256,
                "size_bytes": self.size_bytes, "truncated": self.truncated,
                "tables": [t.as_json() for t in self.tables], "metadata": self.metadata,
                "notes": self.notes, "error": self.error}

    def as_markdown(self) -> str:
        """The whole document as one markdown file — what gets written to data/documents/ and
        what a model is later shown."""
        head = [f"# {self.title}", "", f"*Source: `{self.path.name}` · {self.describe()}*"]
        for key in ("author", "created", "modified", "subject"):
            if self.metadata.get(key):
                head.append(f"*{key.capitalize()}: {self.metadata[key]}*")
        for note in self.notes:
            head.append(f"*Note: {note}*")
        return "\n".join(head) + "\n\n" + self.text


# ── helpers ──────────────────────────────────────────────────────────────────
def sha256_file(path: Path, chunk: int = 1 << 20) -> str:
    h = hashlib.sha256()
    with open(path, "rb") as fh:
        while block := fh.read(chunk):
            h.update(block)
    return h.hexdigest()


def _cell(value: Any) -> str:
    if value is None:
        return ""
    text = str(value).replace("\r", " ").replace("\n", " ").replace("|", "\\|").strip()
    return text[:MAX_CELL_CHARS]


def _clip_rows(rows: list[list[str]]) -> tuple[list[list[str]], bool]:
    clipped = [r[:MAX_TABLE_COLS] for r in rows[:MAX_TABLE_ROWS]]
    return clipped, len(rows) > MAX_TABLE_ROWS


class _Budget:
    """Wall-clock guard, checked between pages/sheets/slides."""

    def __init__(self, seconds: float) -> None:
        self.deadline = time.monotonic() + seconds
        self.seconds = seconds

    def spent(self) -> bool:
        return time.monotonic() > self.deadline


def _zip_expanded_size(path: Path) -> int:
    with zipfile.ZipFile(path) as zf:
        return sum(i.file_size for i in zf.infolist())


def _zip_member_text(path: Path, name: str) -> str:
    try:
        with zipfile.ZipFile(path) as zf:
            return zf.read(name).decode("utf-8", "replace")
    except (KeyError, OSError, zipfile.BadZipFile):
        return ""


def _read_text_bytes(raw: bytes) -> str:
    for encoding in ("utf-8-sig", "utf-8", "utf-16", "cp1252", "latin-1"):
        try:
            return raw.decode(encoding)
        except (UnicodeDecodeError, UnicodeError):
            continue
    return raw.decode("utf-8", "replace")


# ── per-format readers ───────────────────────────────────────────────────────
def _read_pdf(path: Path, doc: Document, budget: _Budget) -> None:
    from pypdf import PdfReader

    # pypdf logs "EOF marker not found" and friends straight to stderr. RevenueOS says why a
    # document could not be read in one sentence of its own; the library's commentary is noise.
    pypdf_log = logging.getLogger("pypdf")
    previous, pypdf_log.level = pypdf_log.level, logging.CRITICAL
    try:
        _read_pdf_inner(path, doc, budget, PdfReader)
    finally:
        pypdf_log.level = previous


def _read_pdf_inner(path: Path, doc: Document, budget: _Budget, PdfReader: Any) -> None:
    reader = PdfReader(str(path), strict=False)
    if reader.is_encrypted:
        try:
            opened = reader.decrypt("")
        except Exception:
            opened = 0
        if not opened:
            doc.error = "this PDF is password-protected. Remove the password (open it and re-save without one) and drop it in again."
            return
    total = len(reader.pages)
    doc.pages, doc.unit = total, "page"
    meta = reader.metadata or {}
    for key, field_name in (("/Title", "title"), ("/Author", "author"), ("/CreationDate", "created"),
                            ("/ModDate", "modified"), ("/Producer", "producer")):
        value = meta.get(key)
        if value:
            doc.metadata[field_name] = str(value)[:300]
    limit = min(total, MAX_UNITS)
    if total > MAX_UNITS:
        doc.notes.append(f"read the first {MAX_UNITS} of {total} pages (the page cap)")
        doc.truncated = True
    parts: list[str] = []
    read = 0
    for i in range(limit):
        if budget.spent():
            doc.notes.append(f"stopped after {read} of {total} pages: the {budget.seconds:.0f}s extraction budget ran out")
            doc.truncated = True
            break
        try:
            page_text = reader.pages[i].extract_text() or ""
        except Exception as exc:  # one broken page must not lose the other 200
            doc.notes.append(f"page {i + 1} could not be read ({type(exc).__name__})")
            page_text = ""
        read += 1
        if page_text.strip():
            parts.append(f"## Page {i + 1}\n\n{page_text.strip()}")
    doc.text = "\n\n".join(parts)
    if not doc.text.strip():
        doc.error = SCAN_SENTENCE
        return
    empty = read - len(parts)
    if empty:
        doc.notes.append(f"{empty} of the {read} pages read carry no text layer (scanned or image-only pages)")
    doc.notes.append("PDF tables are not extracted as rows: the format stores them as positioned text, not as a table.")
    doc.ok = True


def _docx_page_count(path: Path) -> int | None:
    """Word stores a page count in docProps/app.xml when *it* saved the file; a .docx written by
    a library has none. Never guessed."""
    m = re.search(r"<Pages>(\d+)</Pages>", _zip_member_text(path, "docProps/app.xml"))
    return int(m.group(1)) if m else None


def _docx_blocks(d: Any) -> Any:
    """Paragraphs and tables in the order they appear in the document. python-docx exposes them
    as two separate lists (`.paragraphs`, `.tables`), which would move every table to the end and
    show the reader a document the business never wrote; the body's own child order is the truth."""
    import docx.table
    import docx.text.paragraph

    body = d.element.body
    for child in body.iterchildren():
        if child.tag.endswith("}p"):
            yield docx.text.paragraph.Paragraph(child, d)
        elif child.tag.endswith("}tbl"):
            yield docx.table.Table(child, d)


def _read_docx(path: Path, doc: Document, budget: _Budget) -> None:
    import docx
    import docx.table

    d = docx.Document(str(path))
    core = d.core_properties
    for value, field_name in ((core.title, "title"), (core.author, "author"), (core.created, "created"),
                              (core.modified, "modified"), (core.subject, "subject")):
        if value:
            doc.metadata[field_name] = str(value)[:300]
    doc.pages = _docx_page_count(path)
    doc.unit = "page" if doc.pages is not None else None
    if doc.pages is None:
        doc.notes.append("no page count: a .docx only carries one when the application that saved it wrote one")
    else:
        doc.notes.append("the page count is the one the application that last saved this .docx recorded; "
                         "RevenueOS does not lay the document out to count pages itself")
    parts: list[str] = []
    for block in _docx_blocks(d):
        if budget.spent():
            doc.notes.append(f"stopped part-way through the document: the {budget.seconds:.0f}s extraction budget ran out")
            doc.truncated = True
            break
        if isinstance(block, docx.table.Table):
            rows, clipped = _clip_rows([[_cell(c.text) for c in row.cells] for row in block.rows])
            if not rows:
                continue
            t = Table(name=f"Table {len(doc.tables) + 1}", rows=rows, truncated=clipped)
            doc.tables.append(t)
            parts.append(f"{t.as_markdown()}")
            continue
        text = block.text.strip()
        if not text:
            continue
        style = (block.style.name if block.style is not None else "") or ""
        m = re.match(r"Heading (\d)", style)
        if m:
            parts.append("#" * min(int(m.group(1)), 6) + " " + text)
        elif style in ("Title", "Subtitle"):
            parts.append("# " + text)
        elif style.startswith("List"):
            parts.append("- " + text)
        else:
            parts.append(text)
    doc.text = "\n\n".join(parts)
    if not doc.text.strip():
        doc.error = "this .docx has no text in it (no paragraphs and no tables)."
        return
    doc.ok = True


def _read_xlsx(path: Path, doc: Document, budget: _Budget) -> None:
    import openpyxl

    wb = openpyxl.load_workbook(str(path), read_only=True, data_only=True)
    try:
        names = wb.sheetnames
        doc.pages, doc.unit = len(names), "sheet"
        props = wb.properties
        for value, field_name in ((props.title, "title"), (props.creator, "author"),
                                  (props.created, "created"), (props.modified, "modified")):
            if value:
                doc.metadata[field_name] = str(value)[:300]
        parts: list[str] = []
        for name in names[:MAX_UNITS]:
            if budget.spent():
                doc.notes.append(f"stopped after {len(doc.tables)} of {len(names)} sheets: "
                                 f"the {budget.seconds:.0f}s extraction budget ran out")
                doc.truncated = True
                break
            ws = wb[name]
            rows: list[list[str]] = []
            for row in ws.iter_rows(values_only=True):
                cells = [_cell(v) for v in row]
                while cells and cells[-1] == "":
                    cells.pop()
                if cells:
                    rows.append(cells)
                if len(rows) > MAX_TABLE_ROWS:
                    break
            rows, clipped = _clip_rows(rows)
            if not rows:
                parts.append(f"## Sheet: {name}\n\n(empty sheet)")
                continue
            width = max(len(r) for r in rows)
            rows = [(r + [""] * width)[:width] for r in rows]
            t = Table(name=name, rows=rows, truncated=clipped)
            doc.tables.append(t)
            parts.append(f"## Sheet: {name}\n\n{t.as_markdown()}")
        if len(names) > MAX_UNITS:
            doc.notes.append(f"read the first {MAX_UNITS} of {len(names)} sheets (the sheet cap)")
            doc.truncated = True
        doc.text = "\n\n".join(parts)
    finally:
        wb.close()
    if not doc.tables:
        doc.error = "this spreadsheet has no data in any sheet."
        return
    doc.ok = True


def _read_csv(path: Path, doc: Document, budget: _Budget) -> None:
    raw = path.read_bytes()
    text = _read_text_bytes(raw)
    sample = text[:8192]
    try:
        dialect: Any = csv.Sniffer().sniff(sample, delimiters=",;\t|")
    except csv.Error:
        dialect = csv.excel_tab if path.suffix.lower() == ".tsv" else csv.excel
    rows = [[_cell(c) for c in row] for row in csv.reader(io.StringIO(text), dialect) if any(str(c).strip() for c in row)]
    if not rows:
        doc.error = "this CSV file has no rows in it."
        return
    width = max(len(r) for r in rows)
    clipped_rows, clipped = _clip_rows([(r + [""] * width)[:width] for r in rows])
    t = Table(name=path.stem, rows=clipped_rows, truncated=clipped)
    doc.tables.append(t)
    doc.text = f"## {path.stem}\n\n{t.as_markdown()}"
    doc.notes.append(f"{len(rows)} row(s), {width} column(s)")
    doc.ok = True


def _read_pptx(path: Path, doc: Document, budget: _Budget) -> None:
    from pptx import Presentation

    prs = Presentation(str(path))
    core = prs.core_properties
    for value, field_name in ((core.title, "title"), (core.author, "author"), (core.created, "created"),
                              (core.modified, "modified"), (core.subject, "subject")):
        if value:
            doc.metadata[field_name] = str(value)[:300]
    slides = list(prs.slides)
    doc.pages, doc.unit = len(slides), "slide"
    parts: list[str] = []
    for i, slide in enumerate(slides[:MAX_UNITS], start=1):
        if budget.spent():
            doc.notes.append(f"stopped after {i - 1} of {len(slides)} slides: the {budget.seconds:.0f}s extraction budget ran out")
            doc.truncated = True
            break
        title = ""
        try:
            if slide.shapes.title is not None:
                title = (slide.shapes.title.text or "").strip()
        except (AttributeError, ValueError):
            title = ""
        lines: list[str] = [f"## Slide {i}" + (f": {title}" if title else "")]
        for shape in slide.shapes:
            if shape.has_text_frame:
                body = (shape.text_frame.text or "").strip()
                if body and body != title:
                    lines.append(body)
            if getattr(shape, "has_table", False):
                rows = [[_cell(c.text) for c in row.cells] for row in shape.table.rows]
                rows, clipped = _clip_rows(rows)
                if rows:
                    t = Table(name=f"Slide {i} table", rows=rows, truncated=clipped)
                    doc.tables.append(t)
                    lines.append(t.as_markdown())
        if slide.has_notes_slide:
            notes_text = (slide.notes_slide.notes_text_frame.text or "").strip()
            if notes_text:
                lines.append(f"*Speaker notes:* {notes_text}")
        parts.append("\n\n".join(lines))
    if len(slides) > MAX_UNITS:
        doc.notes.append(f"read the first {MAX_UNITS} of {len(slides)} slides (the slide cap)")
        doc.truncated = True
    doc.text = "\n\n".join(parts)
    if not re.sub(r"(?m)^## Slide \d+.*$", "", doc.text).strip():
        doc.error = "this deck has no text on any slide (the slides may be images)."
        return
    doc.ok = True


def _read_plain(path: Path, doc: Document, budget: _Budget) -> None:
    doc.text = _read_text_bytes(path.read_bytes()).strip()
    if not doc.text:
        doc.error = "this file is empty."
        return
    doc.ok = True


def _read_rtf(path: Path, doc: Document, budget: _Budget) -> None:
    from striprtf.striprtf import rtf_to_text

    raw = _read_text_bytes(path.read_bytes())
    if not raw.lstrip().startswith("{\\rtf"):
        doc.error = "this file is named .rtf but does not start with an RTF header, so it is not an RTF document."
        return
    doc.text = rtf_to_text(raw, errors="ignore").strip()
    if not doc.text:
        doc.error = "this RTF document has no text in it."
        return
    doc.ok = True


_READERS = {"pdf": _read_pdf, "docx": _read_docx, "xlsx": _read_xlsx, "csv": _read_csv,
            "pptx": _read_pptx, "text": _read_plain, "rtf": _read_rtf}
_ZIP_FORMATS = {"docx", "xlsx", "pptx"}
_MISSING_LIB = {"pdf": "pypdf", "docx": "python-docx", "xlsx": "openpyxl", "pptx": "python-pptx", "rtf": "striprtf"}


# ── the one entry point ──────────────────────────────────────────────────────
def read_document(path: Path | str, *, max_bytes: int = MAX_BYTES, max_chars: int = MAX_CHARS,
                  time_budget: float = TIME_BUDGET_S) -> Document:
    """Read one business document. Never raises: an unreadable file comes back `ok=False` with
    one plain sentence in `error`."""
    path = Path(path)
    ext = path.suffix.lower()
    fmt, _unit = FORMATS.get(ext, ("unknown", None))
    doc = Document(path=path, format=fmt)

    if not path.is_file():
        doc.error = "there is no file at that path."
        return doc
    doc.size_bytes = path.stat().st_size
    if ext in REFUSED:
        doc.error = f"{ext} is {REFUSED[ext]}"
        return doc
    if fmt == "unknown":
        doc.error = (f"RevenueOS cannot read {ext or 'a file with no extension'} yet. "
                     f"It reads {', '.join(sorted(SUPPORTED_EXTENSIONS))}.")
        return doc
    if doc.size_bytes == 0:
        doc.error = "this file is empty (0 bytes)."
        return doc
    if doc.size_bytes > max_bytes:
        doc.error = (f"this file is {doc.size_bytes / 1048576:.1f} MB and the limit is "
                     f"{max_bytes / 1048576:.0f} MB. Split it, or export the part that matters.")
        return doc
    try:
        doc.sha256 = sha256_file(path)
    except OSError as exc:
        doc.error = f"this file could not be read from disk ({exc.strerror or exc})."
        return doc

    if fmt in _ZIP_FORMATS:
        try:
            expanded = _zip_expanded_size(path)
        except zipfile.BadZipFile:
            doc.error = f"this file is named {ext} but is not a valid Office document (it is not a zip archive)."
            return doc
        except OSError as exc:
            doc.error = f"this file could not be read from disk ({exc.strerror or exc})."
            return doc
        if expanded > MAX_UNCOMPRESSED:
            doc.error = (f"this {ext} expands to {expanded / 1048576:.0f} MB, over the "
                         f"{MAX_UNCOMPRESSED / 1048576:.0f} MB limit; RevenueOS will not open it.")
            return doc

    budget = _Budget(time_budget)
    try:
        _READERS[fmt](path, doc, budget)
    except ImportError:
        doc.ok = False
        doc.error = (f"the library that reads {ext} is not installed in this environment "
                     f"(pip install {_MISSING_LIB.get(fmt, fmt)}).")
        return doc
    except Exception as exc:  # every parser failure becomes one sentence, never a traceback
        doc.ok = False
        doc.text = ""
        doc.tables = []
        detail = " ".join(str(exc).split())[:200]
        doc.error = f"this {ext} could not be read: {type(exc).__name__}{': ' + detail if detail else ''}."
        return doc

    if len(doc.text) > max_chars:
        doc.text = doc.text[:max_chars].rstrip() + TRUNCATION_MARKER.format(limit=max_chars)
        doc.truncated = True
        doc.notes.append(f"text truncated at {max_chars:,} characters")
    if not doc.ok and not doc.error:  # defensive: a reader must always say why
        doc.error = f"this {ext} could not be read."
    return doc


def inbox_documents(directory: Path) -> list[Path]:
    """Every file in the drop folder RevenueOS knows how to try — including the ones it will
    refuse by name, so the customer is told rather than silently ignored."""
    if not directory.is_dir():
        return []
    known = set(SUPPORTED_EXTENSIONS) | set(REFUSED)
    return sorted(p for p in directory.iterdir()
                  if p.is_file() and not p.name.startswith(".") and p.suffix.lower() in known)
